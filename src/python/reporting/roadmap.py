"""
Remediation Roadmap Word document generation.
Produces *-Microsoft 365 Remediation Roadmap-{date}.docx from Plan.json.
"""

from __future__ import annotations

import logging
from datetime import datetime, timezone
from pathlib import Path

from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

from ..snapshot import get_metadata

log = logging.getLogger(__name__)

_NAVY  = RGBColor(0x1F, 0x38, 0x64)
_SLATE = RGBColor(0x2E, 0x74, 0xB5)

_PHASE_ORDER: dict[str, int] = {"Immediate": 0, "Near Term": 1, "Planned": 2, "Monitor": 3}
_SEV_ORDER:   dict[str, int] = {"High": 0, "Medium": 1, "Low": 2}


# ---------------------------------------------------------------------------
# Public API
# ---------------------------------------------------------------------------

def generate(
    plan: dict,
    snapshot: dict,
    output_path: Path,
) -> None:
    """Write the Remediation Roadmap docx to output_path."""
    output_path = Path(output_path)
    if output_path.suffix.lower() != ".docx":
        output_path = output_path.with_suffix(".docx")
    output_path.parent.mkdir(parents=True, exist_ok=True)

    meta = get_metadata(snapshot)
    tenant_obj = meta.get("Tenant") or {}
    tenant_name = (
        (tenant_obj.get("DisplayName") if isinstance(tenant_obj, dict) else None)
        or meta.get("TenantDisplayName")
        or plan.get("TenantName")
        or "Tenant"
    )
    generated_at = plan.get("GeneratedAt", "")
    if generated_at:
        try:
            dt = datetime.fromisoformat(generated_at.replace("Z", "+00:00"))
            generated_at = dt.strftime("%B %d, %Y")
        except (ValueError, AttributeError):
            pass

    doc = Document()
    _set_default_style(doc)

    _write_cover(doc, tenant_name, generated_at)
    _write_intro(doc, plan, tenant_name)
    _write_roadmap_phases(doc, plan)
    _write_action_detail(doc, plan)

    doc.save(str(output_path))
    log.info("Roadmap written: %s", output_path)


# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------

def _set_default_style(doc: Document) -> None:
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)


def _heading(doc: Document, text: str, level: int, color: RGBColor | None = None) -> None:
    h = doc.add_heading(text, level=level)
    h.style.font.color.rgb = color or _NAVY
    h.paragraph_format.space_after = Pt(6)


def _para(doc: Document, text: str, bold: bool = False, italic: bool = False) -> None:
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.bold = bold
    run.italic = italic
    p.paragraph_format.space_after = Pt(4)


def _label_value(doc: Document, label: str, value: str) -> None:
    p = doc.add_paragraph()
    p.add_run(f"{label}: ").bold = True
    p.add_run(value)
    p.paragraph_format.space_after = Pt(2)


def _write_cover(doc: Document, tenant_name: str, generated_at: str) -> None:
    doc.add_paragraph()
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run(f"{tenant_name}\nMicrosoft 365 Remediation Roadmap")
    run.bold = True
    run.font.size = Pt(22)
    run.font.color.rgb = _NAVY

    doc.add_paragraph()
    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub.add_run(f"Generated: {generated_at}").font.size = Pt(12)

    doc.add_page_break()


def _write_intro(doc: Document, plan: dict, tenant_name: str) -> None:
    _heading(doc, "Introduction", 1)
    findings = plan.get("Findings", [])
    actions = plan.get("RoadmapActions", [])

    immediate = sum(1 for f in findings if f.get("RoadmapPhase") == "Immediate")
    near_term  = sum(1 for f in findings if f.get("RoadmapPhase") == "Near Term")
    planned    = sum(1 for f in findings if f.get("RoadmapPhase") == "Planned")
    monitor    = sum(1 for f in findings if f.get("RoadmapPhase") == "Monitor")

    _para(
        doc,
        f"This roadmap summarizes the prioritized remediation work identified during the "
        f"Microsoft 365 tenant assessment for {tenant_name}. It is organized by execution phase "
        f"and workstream to help the implementation team plan and track remediation efforts.",
    )

    # Summary table
    table = doc.add_table(rows=1, cols=4)
    table.style = "Table Grid"
    hdr = table.rows[0].cells
    for i, h in enumerate(["Phase", "Findings", "Work Items", "Workstreams"]):
        hdr[i].text = h
        hdr[i].paragraphs[0].runs[0].bold = True

    phase_ws: dict[str, set[str]] = {}
    phase_actions: dict[str, int] = {}
    for a in actions:
        ph = a.get("RoadmapPhase", "Monitor")
        phase_ws.setdefault(ph, set()).add(a.get("Workstream", ""))
        phase_actions[ph] = phase_actions.get(ph, 0) + 1

    for phase, count in [
        ("Immediate", immediate), ("Near Term", near_term),
        ("Planned", planned), ("Monitor", monitor),
    ]:
        if count == 0:
            continue
        row = table.add_row().cells
        row[0].text = phase
        row[1].text = str(count)
        row[2].text = str(phase_actions.get(phase, 0))
        row[3].text = ", ".join(sorted(phase_ws.get(phase, set())))

    doc.add_paragraph()
    doc.add_page_break()


def _write_roadmap_phases(doc: Document, plan: dict) -> None:
    """Write a phase-by-phase summary of actions."""
    _heading(doc, "Roadmap Overview", 1)
    actions = sorted(
        plan.get("RoadmapActions", []),
        key=lambda a: (_PHASE_ORDER.get(a.get("RoadmapPhase", "Monitor"), 99), a.get("Workstream", "")),
    )

    current_phase = None
    for action in actions:
        phase = action.get("RoadmapPhase", "Monitor")
        if phase != current_phase:
            current_phase = phase
            _heading(doc, phase, 2, _SLATE)

        # Action summary block
        p = doc.add_paragraph(style="List Bullet")
        run = p.add_run(action.get("ActionTitle", ""))
        run.bold = True
        run = p.add_run(
            f" ({action.get('Workstream', '')} | {action.get('HighestSeverity', '')} | "
            f"{action.get('FindingCount', 0)} finding(s) | {action.get('EstimatedPsHours', '')})"
        )
        run.font.size = Pt(10)

    doc.add_page_break()


def _write_action_detail(doc: Document, plan: dict) -> None:
    """Write full action details with finding tables."""
    _heading(doc, "Action Details", 1)

    actions = sorted(
        plan.get("RoadmapActions", []),
        key=lambda a: (_PHASE_ORDER.get(a.get("RoadmapPhase", "Monitor"), 99), a.get("Workstream", "")),
    )
    findings = plan.get("Findings", [])

    current_phase = None
    for action in actions:
        phase = action.get("RoadmapPhase", "Monitor")
        if phase != current_phase:
            current_phase = phase
            _heading(doc, phase, 2, _SLATE)

        ws = action.get("Workstream", "")
        _heading(doc, action.get("ActionTitle", ws), 3)

        _label_value(doc, "Theme",        action.get("Theme", ""))
        _label_value(doc, "Workstream",   ws)
        _label_value(doc, "Effort",       action.get("EstimatedPsHours", ""))
        _label_value(doc, "Owner",        action.get("PrimaryOwner", ""))
        _label_value(doc, "Quick Win",    "Yes" if action.get("QuickWinEligible") else "No")
        doc.add_paragraph()

        _para(doc, "What This Addresses", bold=True)
        _para(doc, action.get("WhatThisAddresses", ""))

        _para(doc, "Why It Matters", bold=True)
        _para(doc, action.get("WhyItMatters", ""))

        _para(doc, "Recommended Next Step", bold=True)
        _para(doc, action.get("RecommendedNextStep", ""))

        _para(doc, "Business Value", bold=True)
        _para(doc, action.get("BusinessValue", ""))

        _para(doc, "Success Check", bold=True)
        _para(doc, action.get("SuccessCheck", ""))

        _para(doc, "First Validation Step", bold=True)
        _para(doc, action.get("FirstValidationStep", ""))

        # Findings table for this workstream/phase
        ws_findings = [
            f for f in findings
            if f.get("Workstream") == ws and f.get("RoadmapPhase") == phase
        ]
        ws_findings.sort(key=lambda f: _SEV_ORDER.get(f.get("Severity", "Low"), 99))

        if ws_findings:
            _para(doc, f"Findings ({len(ws_findings)})", bold=True)
            tbl = doc.add_table(rows=1, cols=4)
            tbl.style = "Table Grid"
            hdr = tbl.rows[0].cells
            for i, h in enumerate(["Severity", "Reference", "Finding", "Remediation"]):
                hdr[i].text = h
                hdr[i].paragraphs[0].runs[0].bold = True

            for f in ws_findings:
                row = tbl.add_row().cells
                row[0].text = str(f.get("Severity", ""))
                row[1].text = str(f.get("RuleId", ""))
                finding_text = str(f.get("Finding", ""))
                row[2].text = finding_text[:200] + ("..." if len(finding_text) > 200 else "")
                remed = str(f.get("TechnicalRemediation", ""))
                row[3].text = remed[:200] + ("..." if len(remed) > 200 else "")

        doc.add_paragraph()
