"""
Engineer Action Pack generation - Word document format.
Generates EngPack.docx from plan data (produced by analysis.plan.generate).
"""

from __future__ import annotations

import logging
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENTATION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

from ..snapshot import get_metadata

log = logging.getLogger(__name__)

_NAVY   = RGBColor(0x1F, 0x38, 0x64)
_ORANGE = RGBColor(0xC5, 0x5A, 0x11)
_WHITE  = RGBColor(0xFF, 0xFF, 0xFF)

_SEV_BG: dict[str, str] = {
    "High":   "FFE0E0",
    "Medium": "FFF0D8",
    "Low":    "E8F4E8",
}
_SEV_FG: dict[str, RGBColor] = {
    "High":   RGBColor(0xC0, 0x00, 0x00),
    "Medium": RGBColor(0xC5, 0x5A, 0x11),
    "Low":    RGBColor(0x4A, 0x7C, 0x3F),
}

_PHASE_ORDER: dict[str, int] = {"Immediate": 0, "Near Term": 1, "Planned": 2, "Monitor": 3}
_SEV_ORDER:   dict[str, int] = {"High": 0, "Medium": 1, "Low": 2}


def generate(
    plan: dict,
    snapshot: dict,
    output_path: Path,
    snapshot_path: str | None = None,
    support_dir: Path | None = None,
    coverage: dict | None = None,
) -> None:
    """Write EngPack.docx to output_path."""
    output_path = Path(output_path)
    if output_path.suffix.lower() != ".docx":
        output_path = output_path.with_suffix(".docx")
    output_path.parent.mkdir(parents=True, exist_ok=True)

    meta = get_metadata(snapshot)
    tenant_obj = meta.get("Tenant") or {}
    tenant_name = (
        (tenant_obj.get("DisplayName") if isinstance(tenant_obj, dict) else None)
        or meta.get("TenantDisplayName")
        or plan.get("TenantName")
        or "Tenant"
    )
    generated_at = plan.get("GeneratedAt", meta.get("GeneratedAt", ""))

    findings: list[dict] = plan.get("Findings", [])
    findings_sorted = sorted(
        findings,
        key=lambda f: (
            _SEV_ORDER.get(f.get("Severity", "Low"), 99),
            _PHASE_ORDER.get(f.get("RoadmapPhase", "Monitor"), 99),
        ),
    )
    ws_summaries: list[dict] = plan.get("WorkstreamSummaries", [])

    doc = Document()
    _set_landscape(doc)
    _setup_styles(doc)

    _write_cover(doc, tenant_name, generated_at, snapshot_path)
    _write_engineering_summary(doc, findings, ws_summaries)
    _write_findings_table(doc, findings_sorted)
    _write_verification_checklist(doc)
    _write_coverage(doc, coverage)
    _write_supporting_files(doc, support_dir, output_path)

    doc.save(str(output_path))
    log.info("EngPack written: %s (%d findings)", output_path, len(findings))


# ---------------------------------------------------------------------------
# Document setup
# ---------------------------------------------------------------------------

def _set_landscape(doc: Document) -> None:
    section = doc.sections[0]
    new_width = section.page_height
    new_height = section.page_width
    section.page_width  = new_width
    section.page_height = new_height
    section.orientation = WD_ORIENTATION.LANDSCAPE


def _setup_styles(doc: Document) -> None:
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)


def _set_cell_bg(cell, hex_color: str) -> None:
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), hex_color)
    shd.set(qn("w:val"), "clear")
    tcPr.append(shd)


def _heading(doc: Document, text: str, level: int, color: RGBColor | None = None) -> None:
    h = doc.add_heading(text, level=level)
    if h.runs:
        h.runs[0].font.color.rgb = color or _NAVY
    h.paragraph_format.space_after = Pt(6)


# ---------------------------------------------------------------------------
# Cover
# ---------------------------------------------------------------------------

def _write_cover(
    doc: Document,
    tenant_name: str,
    generated_at: str,
    snapshot_path: str | None,
) -> None:
    doc.add_paragraph()
    doc.add_paragraph()

    brand = doc.add_paragraph()
    run = brand.add_run("ARRAYA SOLUTIONS")
    run.font.color.rgb = _ORANGE
    run.font.size = Pt(13)
    run.bold = True

    doc.add_paragraph()

    title = doc.add_paragraph()
    run = title.add_run(tenant_name)
    run.bold = True
    run.font.size = Pt(28)
    run.font.color.rgb = _NAVY

    sub = doc.add_paragraph()
    run = sub.add_run("Engineer Action Pack")
    run.font.size = Pt(18)
    run.font.color.rgb = _NAVY

    doc.add_paragraph()

    meta_p = doc.add_paragraph()
    meta_p.add_run("Generated: ").bold = True
    meta_p.add_run(generated_at)

    if snapshot_path:
        src_p = doc.add_paragraph()
        src_p.add_run("Source snapshot: ").bold = True
        src_p.add_run(snapshot_path)

    doc.add_paragraph()
    sep = doc.add_paragraph()
    run = sep.add_run("_" * 80)
    run.font.color.rgb = _ORANGE

    doc.add_paragraph()
    note = doc.add_paragraph()
    note.add_run(
        "Use this pack to review the current tenant state, confirm each finding, "
        "and plan remediation work that fits your change process."
    )
    note.paragraph_format.space_after = Pt(6)

    doc.add_page_break()


# ---------------------------------------------------------------------------
# Engineering Summary
# ---------------------------------------------------------------------------

def _write_engineering_summary(
    doc: Document,
    findings: list[dict],
    ws_summaries: list[dict],
) -> None:
    _heading(doc, "Engineering Summary", 1)

    phase_counts: dict[str, int] = {p: 0 for p in ("Immediate", "Near Term", "Planned", "Monitor")}
    for f in findings:
        phase = f.get("RoadmapPhase", "Monitor")
        if phase in phase_counts:
            phase_counts[phase] += 1

    summary_data = [
        ("Total Findings", str(len(findings))),
        ("Immediate",      str(phase_counts["Immediate"])),
        ("Near Term",      str(phase_counts["Near Term"])),
        ("Planned",        str(phase_counts["Planned"])),
        ("Monitor",        str(phase_counts["Monitor"])),
    ]

    tbl = doc.add_table(rows=1, cols=len(summary_data))
    tbl.style = "Table Grid"
    for i, (label, value) in enumerate(summary_data):
        cell = tbl.rows[0].cells[i]
        _set_cell_bg(cell, "1F3864")
        p = cell.paragraphs[0]
        p.clear()
        run = p.add_run(label)
        run.bold = True
        run.font.color.rgb = _WHITE
        run.font.size = Pt(9)

    row2 = tbl.add_row().cells
    for i, (label, value) in enumerate(summary_data):
        cell = row2[i]
        p = cell.paragraphs[0]
        run = p.add_run(value)
        run.bold = True
        run.font.size = Pt(18)
        if label == "Immediate":
            run.font.color.rgb = RGBColor(0xC0, 0x00, 0x00)
        elif label == "Near Term":
            run.font.color.rgb = _ORANGE
        elif label == "Total Findings":
            run.font.color.rgb = _NAVY

    doc.add_paragraph()

    if ws_summaries:
        _heading(doc, "Workstream Summary", 2)
        cols = ["Severity", "Workstream", "Area", "Open Findings", "Top Signals"]
        widths = [Inches(0.9), Inches(1.4), Inches(1.4), Inches(1.0), Inches(4.5)]
        tbl2 = doc.add_table(rows=1, cols=len(cols))
        tbl2.style = "Table Grid"
        hdr = tbl2.rows[0].cells
        for i, (label, w) in enumerate(zip(cols, widths)):
            _set_cell_bg(hdr[i], "1F3864")
            hdr[i].width = w
            p = hdr[i].paragraphs[0]
            p.clear()
            run = p.add_run(label)
            run.bold = True
            run.font.color.rgb = _WHITE
            run.font.size = Pt(9)

        for ws in ws_summaries:
            row = tbl2.add_row().cells
            sev = ws.get("Severity", "")
            row[0].text = sev
            row[1].text = ws.get("Workstream", "")
            row[2].text = ws.get("Area", "")
            row[3].text = str(ws.get("OpenFindings", 0))
            row[4].text = ws.get("TopSignals", "")
            if sev in _SEV_BG:
                _set_cell_bg(row[0], _SEV_BG[sev])
                if row[0].paragraphs[0].runs:
                    row[0].paragraphs[0].runs[0].font.color.rgb = _SEV_FG[sev]
                    row[0].paragraphs[0].runs[0].bold = True
            for cell in row:
                for para in cell.paragraphs:
                    for run in para.runs:
                        if not run.font.size:
                            run.font.size = Pt(9)

        doc.add_paragraph()

    _heading(doc, "Before You Start", 2)
    before_items = [
        "Validate tenant admin roles, Graph scopes, Exchange connectivity, "
        "and any pilot exclusions before enforcement changes.",
        "Use the support artifacts for backlog import, validation context, "
        "and change-record attachment as needed.",
        "Pilot security and access-policy changes with a limited scope before broad enforcement.",
    ]
    for item in before_items:
        p = doc.add_paragraph(item, style="List Bullet")
        p.paragraph_format.space_after = Pt(3)

    doc.add_page_break()


# ---------------------------------------------------------------------------
# Findings To Work table
# ---------------------------------------------------------------------------

def _write_findings_table(doc: Document, findings: list[dict]) -> None:
    _heading(doc, "Findings To Work", 1)

    if not findings:
        doc.add_paragraph("No findings identified.")
        return

    cols   = ["Severity", "Phase", "Workstream", "Ref", "Finding", "Technical Remediation", "Done"]
    widths = [Inches(0.8), Inches(0.9), Inches(1.1), Inches(0.8), Inches(2.7), Inches(2.7), Inches(0.5)]

    tbl = doc.add_table(rows=1, cols=len(cols))
    tbl.style = "Table Grid"
    hdr = tbl.rows[0].cells
    for i, (label, w) in enumerate(zip(cols, widths)):
        _set_cell_bg(hdr[i], "1F3864")
        hdr[i].width = w
        p = hdr[i].paragraphs[0]
        p.clear()
        run = p.add_run(label)
        run.bold = True
        run.font.color.rgb = _WHITE
        run.font.size = Pt(9)

    for f in findings:
        sev = f.get("Severity", "")
        row = tbl.add_row().cells
        values = [
            sev,
            f.get("RoadmapPhase", ""),
            f.get("Workstream", ""),
            f.get("RuleId", ""),
            f.get("Finding", ""),
            f.get("TechnicalRemediation", ""),
            "",
        ]
        for i, (val, w) in enumerate(zip(values, widths)):
            row[i].width = w
            p = row[i].paragraphs[0]
            run = p.add_run(str(val))
            run.font.size = Pt(8)
            if i == 0 and sev in _SEV_BG:
                _set_cell_bg(row[i], _SEV_BG[sev])
                run.font.color.rgb = _SEV_FG[sev]
                run.bold = True

    doc.add_page_break()


# ---------------------------------------------------------------------------
# Verification Checklist
# ---------------------------------------------------------------------------

def _write_verification_checklist(doc: Document) -> None:
    _heading(doc, "Verification Checklist", 1)
    items = [
        "Confirm each recommendation against the latest tenant state before making changes.",
        "Pilot security and access-policy changes with a limited scope before broad enforcement.",
        "Re-run M365Collect and Improve after remediation milestones to measure delta and retire closed findings.",
    ]
    for item in items:
        doc.add_paragraph(item, style="List Bullet")
    doc.add_paragraph()


# ---------------------------------------------------------------------------
# Assessment Evidence Coverage
# ---------------------------------------------------------------------------

def _write_coverage(doc: Document, coverage: dict | None) -> None:
    if not coverage:
        return

    _heading(doc, "Assessment Evidence Coverage", 1)

    covered = coverage.get("CoveredObjectiveCount", 0)
    review  = coverage.get("ReviewObjectiveCount", 0)
    total   = coverage.get("ObjectiveCount", 0)

    p = doc.add_paragraph()
    p.add_run(f"{covered} of {total} objectives fully covered").bold = True
    p.add_run(f"  |  {review} objective(s) require review")
    doc.add_paragraph()

    cols   = ["Objective", "Area", "Status", "Confidence", "Datasets"]
    widths = [Inches(1.3), Inches(1.5), Inches(1.0), Inches(1.0), Inches(1.0)]
    tbl = doc.add_table(rows=1, cols=len(cols))
    tbl.style = "Table Grid"
    hdr = tbl.rows[0].cells
    for i, (label, w) in enumerate(zip(cols, widths)):
        _set_cell_bg(hdr[i], "1F3864")
        hdr[i].width = w
        p = hdr[i].paragraphs[0]
        p.clear()
        run = p.add_run(label)
        run.bold = True
        run.font.color.rgb = _WHITE
        run.font.size = Pt(9)

    for obj in coverage.get("Objectives", []):
        present  = obj.get("PresentDatasetCount", 0)
        expected = obj.get("ExpectedDatasetCount", 0)
        row = tbl.add_row().cells
        vals = [
            obj.get("ObjectiveId", ""),
            obj.get("Area", ""),
            obj.get("Status", ""),
            obj.get("Confidence", ""),
            f"{present} of {expected}",
        ]
        for i, val in enumerate(vals):
            row[i].text = str(val)
            if row[i].paragraphs[0].runs:
                row[i].paragraphs[0].runs[0].font.size = Pt(9)

    gaps = coverage.get("CoverageGaps", [])
    if gaps:
        doc.add_paragraph()
        _heading(doc, "Coverage Gaps", 2)
        gcols   = ["Gap", "Area", "Limitation"]
        gwidths = [Inches(1.3), Inches(1.5), Inches(5.0)]
        gtbl = doc.add_table(rows=1, cols=len(gcols))
        gtbl.style = "Table Grid"
        ghdr = gtbl.rows[0].cells
        for i, (label, w) in enumerate(zip(gcols, gwidths)):
            _set_cell_bg(ghdr[i], "1F3864")
            ghdr[i].width = w
            p = ghdr[i].paragraphs[0]
            p.clear()
            run = p.add_run(label)
            run.bold = True
            run.font.color.rgb = _WHITE
            run.font.size = Pt(9)
        for gap in gaps:
            row = gtbl.add_row().cells
            row[0].text = gap.get("GapId", "")
            row[1].text = gap.get("Area", "")
            row[2].text = gap.get("CurrentLimitation", "")
            for cell in row:
                if cell.paragraphs[0].runs:
                    cell.paragraphs[0].runs[0].font.size = Pt(9)

    doc.add_paragraph()


# ---------------------------------------------------------------------------
# Supporting Files
# ---------------------------------------------------------------------------

def _write_supporting_files(
    doc: Document,
    support_dir: Path | None,
    output_path: Path,
) -> None:
    if not support_dir:
        return
    _heading(doc, "Supporting Files", 1)
    stem = output_path.stem.replace("-EngPack", "")
    items = [
        f"Support folder: {support_dir}",
        f"Improvement plan JSON: {support_dir / (stem + '-Plan.json')}",
    ]
    for item in items:
        doc.add_paragraph(item, style="List Bullet")
