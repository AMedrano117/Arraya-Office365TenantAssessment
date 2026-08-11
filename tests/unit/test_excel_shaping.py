"""
Worksheet shaping tests for src/python/reporting/excel.py.

Mirrors the PowerShell coverage in tests/unit/Common.Tests.ps1: container
worksheets must expand to one row per record rather than joining arrays into a
single cell or dropping them.

Run with:  python -m unittest discover -s tests/unit -p "test_*.py"
"""

from __future__ import annotations

import sys
import tempfile
import unittest
from pathlib import Path

_REPO_ROOT = Path(__file__).resolve().parents[2]
if str(_REPO_ROOT) not in sys.path:
    sys.path.insert(0, str(_REPO_ROOT))

from openpyxl import load_workbook  # noqa: E402

from src.python.reporting import excel  # noqa: E402


def _rows(sheet) -> list[dict]:
    values = list(sheet.values)
    if not values:
        return []
    headers = [str(h) if h is not None else "" for h in values[0]]
    return [dict(zip(headers, row)) for row in values[1:]]


class ExcelShapingTests(unittest.TestCase):
    def _export(self, data: dict):
        tmp = Path(tempfile.mkdtemp()) / "shaping.xlsx"
        excel.export(data, tmp)
        return load_workbook(tmp)

    # -- issue 1: AdConnectConfiguration -----------------------------------

    def test_adconnect_container_keeps_sync_services_and_errors(self):
        wb = self._export({
            "AdConnectConfiguration": {
                "Summary": {"OnPremisesSyncEnabled": True, "LastSyncDateTime": "2026-08-10T04:00:00"},
                "SyncServices": [{"ServiceName": "contoso.onmicrosoft.com", "ServerName": "AADC01"}],
                "RecentErrors": [{"ErrorType": "DuplicateAttribute", "ErrorCode": "ATTR-001"}],
                "ErrorCount": 1,
            }
        })
        rows = _rows(wb["AdConnectConfiguration"])
        self.assertEqual(list(rows[0].keys()), ["Section", "Item", "Value", "Notes"])

        sections = {r["Section"] for r in rows}
        self.assertIn("Directory Sync", sections)
        # These were dropped entirely before: only non-list siblings were kept.
        self.assertIn("Sync Services", sections)
        self.assertIn("Recent Errors", sections)
        self.assertTrue(any(r["Item"] == "contoso.onmicrosoft.com" for r in rows))
        self.assertTrue(any(r["Item"] == "DuplicateAttribute" for r in rows))

    # -- issue 2: AuthenticationConfig --------------------------------------

    def test_authentication_config_expands_sso_apps_per_row(self):
        wb = self._export({
            "AuthenticationConfig": {
                "Configuration": {
                    "MFAEnabled": True,
                    "MFAMethods": ["Sms", "MicrosoftAuthenticator"],
                    "SSOApplications": [
                        {"DisplayName": "Contoso CRM"},
                        {"DisplayName": "Contoso HR"},
                        {"DisplayName": "Contoso Wiki"},
                    ],
                }
            }
        })
        rows = _rows(wb["AuthenticationConfig"])
        self.assertEqual(len([r for r in rows if r["Section"] == "SSO Applications"]), 3)
        self.assertEqual(len([r for r in rows if r["Section"] == "MFA Methods"]), 2)
        self.assertTrue(any(r["Item"] == "Contoso Wiki" for r in rows))

    def test_flat_authentication_config_is_left_alone(self):
        """The Python collector emits a flat scalar dict; it should stay one row."""
        wb = self._export({
            "AuthenticationConfig": {
                "SecurityDefaultsEnabled": True,
                "AuthenticationMethodPolicy": "authenticationMethodsPolicy",
                "CAPolicyCount": 12,
            }
        })
        rows = _rows(wb["AuthenticationConfig"])
        self.assertEqual(len(rows), 1)
        self.assertIn("CAPolicyCount", rows[0])

    # -- issue 3: MfaEnrollmentSummary --------------------------------------

    def test_mfa_summary_expands_one_row_per_method_python_shape(self):
        wb = self._export({
            "MfaEnrollmentSummary": {
                "TotalUsers": 10,
                "RegisteredUsers": 8,
                "MethodCounts": {"microsoftAuthenticator": 6, "sms": 3, "fido2": 1},
            }
        })
        rows = _rows(wb["MfaEnrollmentSummary"])
        self.assertEqual(list(rows[0].keys()),
                         ["Category", "Method", "UserCount", "PercentOfUsers", "Notes"])
        registered = [r for r in rows if r["Category"] == "Registered"]
        self.assertEqual(len(registered), 3)
        # Graph enum values stay verbatim.
        self.assertTrue(any(r["Method"] == "microsoftAuthenticator" for r in registered))
        sms = next(r for r in registered if r["Method"] == "sms")
        self.assertEqual(sms["PercentOfUsers"], 30.0)

    def test_mfa_summary_reads_count_maps_from_powershell_snapshot(self):
        """A PowerShell snapshot pre-joins the breakdown and keeps raw maps elsewhere."""
        wb = self._export({
            "MfaEnrollmentSummary": {
                "TotalUsers": 10,
                "RegisteredMethodBreakdown": "microsoftAuthenticator (6); sms (3)",
            },
            "MfaRegistrationSummary": {
                "MethodCounts": {"microsoftAuthenticator": 6, "sms": 3},
                "WeakMethodCounts": {"sms": 3},
            },
        })
        rows = _rows(wb["MfaEnrollmentSummary"])
        self.assertEqual(len([r for r in rows if r["Category"] == "Registered"]), 2)
        self.assertEqual(len([r for r in rows if r["Category"] == "Weak"]), 1)

    # -- issue 5: TeamsVoice -------------------------------------------------

    def test_teamsvoice_container_expands_per_record(self):
        wb = self._export({
            "TeamsVoice": {
                "Summary": {"VoiceUserCount": 3, "DataSource": "TeamsPowerShell"},
                "CallingPolicies": [{"Identity": "Global"}, {"Identity": "Tag:NoPSTN"}],
                "PhoneNumbers": [{"TelephoneNumber": "+15555550100"}],
            }
        })
        self.assertIn("TeamsVoice", wb.sheetnames)
        rows = _rows(wb["TeamsVoice"])
        self.assertEqual(len([r for r in rows if r["Section"] == "Calling Policies"]), 2)
        # A leading '+' must survive as text.
        self.assertTrue(any(str(r["Item"]) == "+15555550100" for r in rows))

    # -- ragged rows ---------------------------------------------------------

    def test_ragged_records_keep_later_columns(self):
        wb = self._export({
            "BestPracticeFindings": [
                {"RuleId": "ID-001", "Severity": "High"},
                {"RuleId": "ID-002", "Severity": "Low", "Recommendation": "Enable CA policy"},
            ]
        })
        rows = _rows(wb["BestPracticeFindings"])
        self.assertIn("Recommendation", rows[0])
        second = next(r for r in rows if r["RuleId"] == "ID-002")
        self.assertEqual(second["Recommendation"], "Enable CA policy")

    # -- regression guards ---------------------------------------------------

    def test_uniform_lookup_table_is_not_reshaped(self):
        wb = self._export({
            "Users": {
                "ana@contoso.com": {"DisplayName": "Ana", "UserPrincipalName": "ana@contoso.com"},
                "bo@contoso.com": {"DisplayName": "Bo", "UserPrincipalName": "bo@contoso.com"},
            }
        })
        rows = _rows(wb["Users"])
        self.assertEqual(len(rows), 2)
        self.assertIn("DisplayName", rows[0])
        self.assertNotIn("Section", rows[0])

    def test_single_summary_container_stays_one_row(self):
        wb = self._export({
            "CollaborationActivitySummary": {
                "Summary": {"PeriodDuration": "D180", "TeamsActiveUsers": 7},
            }
        })
        rows = _rows(wb["CollaborationActivitySummary"])
        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0]["PeriodDuration"], "D180")

    def test_label_preserves_acronyms(self):
        self.assertEqual(excel._label("SyncServices"), "Sync Services")
        self.assertEqual(excel._label("SSOApplications"), "SSO Applications")
        self.assertEqual(excel._label("MFAMethods"), "MFA Methods")
        self.assertEqual(excel._label("OnPremisesSyncEnabled"), "On Premises Sync Enabled")


if __name__ == "__main__":
    unittest.main()
